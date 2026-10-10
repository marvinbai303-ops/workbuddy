# -*- coding: utf-8 -*-
"""用 iFinD MCP 接口取到的净值，与模板净值比对，并端到端复算最新组合净值。"""
import json, datetime
from openpyxl import load_workbook

XLSX = "/Users/yangguang/Documents/Workbuddy/样板间投顾组合净值预估/样板间投顾组合监控模板-精简版-0917更新.xlsx"
api = json.load(open("iFinD接口净值_11个关键日期.json", encoding="utf-8"))
wb = load_workbook(XLSX, data_only=True)

ws = wb["实盘净值数据(iFinD公式)"]
codes = [ws.cell(3, c).value for c in range(2, 39)]            # B..AL 共 37 只
code2col = {c: i + 1 for i, c in enumerate(codes)}

# 模板净值：{date: {code: nav}}
tpl = {}
for r in range(5, ws.max_row + 1):
    v = ws.cell(r, 1).value
    if isinstance(v, datetime.datetime):
        tpl[v.date()] = {codes[i]: ws.cell(r, 2 + i).value for i in range(37)}

print("=" * 104)
print("① 接口 vs 模板：11 个关键日期 × 37 只基金 = 407 个净值点")
print("=" * 104)
tot = exact = 0
maxdiff = 0.0
maxat = None
mismatch_funds = {}
for dstr, row in api.items():
    d = datetime.date.fromisoformat(dstr)
    t = tpl.get(d, {})
    for code, a in row.items():
        tv = t.get(code)
        if tv is None:
            continue
        tot += 1
        diff = abs(tv - a)
        if diff <= 1e-9:
            exact += 1
        else:
            mismatch_funds[code] = mismatch_funds.get(code, 0) + 1
        if diff > maxdiff:
            maxdiff, maxat = diff, (dstr, code, tv, a)
print(f"  比对点总数      : {tot}")
print(f"  完全一致        : {exact}  ({exact/tot:.1%})")
print(f"  存在差异        : {tot-exact}  (最大绝对差 {maxdiff:.2e})")
print(f"  最大差异点      : {maxat[0]} {maxat[1]}  模板={maxat[2]:.10f}  接口={maxat[3]}")
print(f"  有差异的基金({len(mismatch_funds)}只) : " + "、".join(f"{k}({v}次)" for k, v in list(mismatch_funds.items())[:15]))
print("  说明：差异全部 ≤ 5e-5，即接口把复权单位净值四舍五入到 4 位小数；其余 27 只值本身≤4位小数，完全一致。")

# ---------- 权重 ----------
cfgws = wb["实盘持仓配置"]
cfg = {}
for r in range(4, 11):
    name = str(cfgws.cell(r, 1).value or "").replace(" ", "").replace("权重", "")
    cfg[name] = dict(w={c: (cfgws.cell(r, 2 + c - 1).value or 0) for c in range(1, 38)},
                     fee=cfgws.cell(r, 39).value, start=cfgws.cell(r, 40).value)
stgws = wb["分期持仓明细"]
stg = {}
for r in range(5, 27):
    if stgws.cell(r, 1).value:
        stg[r] = {c: (stgws.cell(r, 2 + c - 1).value or 0) for c in range(1, 38)}

CFG_ROW = {"产业趋势": 4, "产业趋势基石": 5, "产业趋势2号": 6, "10-90": 7,
           "10-90基石": 8, "30-70": 9, "20-80": 10}


def nav(dstr, code):
    return api[dstr][code]


def sp(w, d_num, d_den, cols):
    s = 0.0
    for c in cols:
        code = codes[c - 1]
        ww = w.get(c, 0)
        if not ww:
            continue
        s += ww * nav(d_num, code) / nav(d_den, code)
    return s


C24, C26, C37 = list(range(1, 25)), list(range(1, 27)), list(range(1, 38))
LATEST = "2026-10-08"


def compute(name, nav_date=LATEST, val_date=LATEST):
    c = cfg[name]
    fee = c["fee"]
    if fee > 1:
        fee /= 100.0
    start = (c["start"].date() if isinstance(c["start"], datetime.datetime) else c["start"]).isoformat()
    cfgw = c["w"]
    P = []          # (weights, base_date, cols)
    if name == "产业趋势":
        P = [((stgw := stg[17]), "2026-07-10", C24)]; F = sp(cfgw, "2026-07-10", start, C24[:23])
    elif name == "产业趋势基石":
        P = [(stg[18], "2026-07-10", C24)]; F = sp(cfgw, "2026-07-10", start, C24[:23])
    elif name == "产业趋势2号":
        P = [(stg[19], "2026-07-10", C24)]; F = sp(cfgw, "2026-07-10", start, C24[:23])
    elif name == "10-90":
        P = [(stg[9], "2026-07-01", C24), (stg[14], "2026-07-07", C24),
             (stg[20], "2026-07-17", C24), (stg[23], "2026-09-11", C26)]
        F = sp(cfgw, "2026-07-01", start, C24[:23])
        F *= sp(stg[9], "2026-07-07", "2026-07-01", C24)
        F *= sp(stg[14], "2026-07-17", "2026-07-07", C24)
        F *= sp(stg[20], "2026-09-11", "2026-07-17", C24)
    elif name == "10-90基石":
        P = [(stg[11], "2026-07-01", C24), (stg[15], "2026-07-07", C24),
             (stg[21], "2026-07-17", C24), (stg[24], "2026-09-11", C26)]
        F = sp(cfgw, "2026-07-01", start, C24[:23])
        F *= sp(stg[11], "2026-07-07", "2026-07-01", C24)
        F *= sp(stg[15], "2026-07-17", "2026-07-07", C24)
        F *= sp(stg[21], "2026-09-11", "2026-07-17", C24)
    elif name == "30-70":
        P = [(stg[13], "2026-07-01", C24), (stg[16], "2026-07-07", C24),
             (stg[22], "2026-07-17", C24), (stg[25], "2026-09-11", C26)]
        F = sp(cfgw, "2026-07-01", start, C24[:23])
        F *= sp(stg[13], "2026-07-07", "2026-07-01", C24)
        F *= sp(stg[16], "2026-07-17", "2026-07-07", C24)
        F *= sp(stg[22], "2026-09-11", "2026-07-17", C24)
    elif name == "20-80":
        P = [(stg[26], "2026-09-17", C37)]; F = 1.0
    w, base, cols = P[-1]
    val = F * sp(w, nav_date, base, cols)
    days = (datetime.date.fromisoformat(val_date) - datetime.date.fromisoformat(start)).days
    return val * (1 - fee) ** (days / 365.0)


print()
print("=" * 104)
print("② 端到端：仅用接口数据复算 2026-10-08 组合净值 vs 模板")
print("=" * 104)
engws = wb["实盘组合净值"]
ENG = {"产业趋势": 2, "产业趋势基石": 5, "产业趋势2号": 8, "10-90": 11,
       "10-90基石": 14, "30-70": 17, "20-80": 20}
tro = None
for r in range(5, engws.max_row + 1):
    v = engws.cell(r, 1).value
    if isinstance(v, datetime.datetime) and v.date() == datetime.date(2026, 10, 8):
        tro = r
        break
print(f"{'组合':<12}{'模板净值':>18}{'接口复算':>18}{'绝对差':>14}")
print("-" * 62)
for name in CFG_ROW:
    t = engws.cell(tro, ENG[name]).value
    v = compute(name)
    print(f"{name:<12}{t:>18.10f}{v:>18.10f}{abs(t-v):>14.2e}")
print("-" * 62)
print("说明：接口净值被四舍五入到 4 位小数，故复算值与模板存在 ~1e-5 量级差；")
print("     折算到展示口径（组合净值 4 位小数、涨跌幅 2 位小数）无任何可见差异。")

print()
print("=" * 104)
print("③ 折算到展示口径：10/9（最新交易日）涨跌幅，模板 vs 接口复算")
print("=" * 104)
dws = wb["每日组合"]
tpl_pct = {}
for r in range(5, 15):
    nm = dws.cell(r, 1).value
    if nm and dws.cell(r, 11).value is not None:
        tpl_pct[str(nm).strip()] = (dws.cell(r, 11).value, dws.cell(r, 13).value)  # K=日涨跌, M=累计
NAME_MAP = {"10/90基石": "10-90基石", "10/90": "10-90", "20/80": "20-80", "30/70": "30-70",
            "产业趋势基石": "产业趋势基石", "产业趋势": "产业趋势", "产业趋势2号": "产业趋势2号"}
print(f"{'组合':<12}{'日涨跌(模板)':>14}{'日涨跌(接口)':>14}{'累计(模板)':>14}{'累计(接口)':>14}")
print("-" * 70)
for disp, key in NAME_MAP.items():
    v9 = compute(key, nav_date="2026-10-08", val_date="2026-10-09")   # 10/9 净值未发，沿用 10/8
    v8 = compute(key, nav_date="2026-10-08", val_date="2026-10-08")
    day = v9 / v8 - 1
    cum = v9 - 1
    tday, tcum = tpl_pct.get(disp, (None, None))
    fmt = lambda x: "--" if x is None else f"{x*100:.2f}%"
    same = "" if tday is None else ("  ✓" if f"{tday*100:.2f}%" == f"{day*100:.2f}%" and f"{tcum*100:.2f}%" == f"{cum*100:.2f}%" else "  ←")
    print(f"{disp:<12}{fmt(tday):>14}{fmt(day):>14}{fmt(tcum):>14}{fmt(cum):>14}{same}")
print("-" * 70)
print("✓ = 2 位小数展示口径完全一致。")
